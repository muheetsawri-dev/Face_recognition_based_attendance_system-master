import os
import json
import datetime
from pathlib import Path

import cv2
import numpy as np
import pandas as pd
import streamlit as st
from PIL import Image
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
STUDENT_DIR = BASE_DIR / "StudentDetails"
TRAIN_DIR = BASE_DIR / "TrainingImage"
LABEL_DIR = BASE_DIR / "TrainingImageLabel"
ATT_DIR = BASE_DIR / "Attendance"
STUDENT_XLSX = STUDENT_DIR / "StudentDetails.xlsx"
MODEL = LABEL_DIR / "Trainner.yml"
SETTINGS_FILE = LABEL_DIR / "settings.json"
CASCADE = BASE_DIR / "haarcascade_frontalface_default.xml"

for p in (STUDENT_DIR, TRAIN_DIR, LABEL_DIR, ATT_DIR):
    p.mkdir(exist_ok=True)

DEFAULT = {"class_start_time": "09:00", "confidence_threshold": 50}

def settings():
    try:
        return {**DEFAULT, **json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))}
    except Exception:
        SETTINGS_FILE.write_text(json.dumps(DEFAULT, indent=2), encoding="utf-8")
        return DEFAULT.copy()

SET = settings()

def ensure_students():
    if not STUDENT_XLSX.exists():
        wb = Workbook()
        ws = wb.active
        ws.append(["SERIAL NO.", "ID", "NAME"])
        wb.save(STUDENT_XLSX)
        wb.close()

def students():
    ensure_students()
    wb = load_workbook(STUDENT_XLSX, read_only=True)
    ws = wb.active
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0] is not None:
            rows.append((int(r[0]), str(r[1]), str(r[2])))
    wb.close()
    return rows

def add_student(sid, name):
    wb = load_workbook(STUDENT_XLSX)
    wb.active.append([next_serial(), sid, name])
    wb.save(STUDENT_XLSX)
    wb.close()

def next_serial():
    rows = students()
    return max((r[0] for r in rows), default=0) + 1

def cascade():
    if not CASCADE.exists():
        return None
    c = cv2.CascadeClassifier(str(CASCADE))
    return c if not c.empty() else None

def faces_from_bytes(data):
    image = Image.open(data).convert("RGB")
    frame = np.asarray(image)
    gray = cv2.cvtColor(frame, cv2.COLOR_RGB2GRAY)
    c = cascade()
    if c is None:
        return gray, []
    return gray, c.detectMultiScale(gray, 1.2, 5, minSize=(60, 60))

def save_sample(serial, sid, name, data, number):
    gray, faces = faces_from_bytes(data)
    if len(faces) == 0:
        return False, "No face detected. Capture a clear front-facing photo."
    x, y, w, h = max(faces, key=lambda f: f[2] * f[3])
    safe = "".join(ch if ch.isalnum() or ch in " _-" else "_" for ch in name).strip()
    path = TRAIN_DIR / f"{safe}.{serial}.{sid}.{number}.jpg"
    cv2.imwrite(str(path), gray[y:y+h, x:x+w])
    return True, str(path)

def train():
    if not hasattr(cv2, "face"):
        return False, "Install opencv-contrib-python."
    images, labels = [], []
    for f in TRAIN_DIR.glob("*.jpg"):
        parts = f.name.split(".")
        if len(parts) < 3:
            continue
        try:
            label = int(parts[1])
        except ValueError:
            continue
        img = cv2.imread(str(f), cv2.IMREAD_GRAYSCALE)
        if img is not None:
            images.append(img)
            labels.append(label)
    if not images:
        return False, "No training images found."
    rec = cv2.face.LBPHFaceRecognizer_create()
    rec.train(images, np.array(labels, dtype=np.int32))
    rec.write(str(MODEL))
    return True, f"Model trained with {len(images)} images."

def attendance_file(date=None):
    date = date or datetime.date.today().strftime("%d-%m-%Y")
    return ATT_DIR / f"Attendance_{date}.xlsx"

def attendance(date=None):
    p = attendance_file(date)
    if not p.exists():
        return []
    wb = load_workbook(p, read_only=True)
    ws = wb.active
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r and r[0] is not None:
            v = list(r) + [None] * 5
            out.append({"Id": str(v[0]), "Name": str(v[1]), "Date": str(v[2]),
                        "In-Time": str(v[3]), "Status": str(v[4])})
    wb.close()
    return out

def mark(sid, name):
    date = datetime.date.today().strftime("%d-%m-%Y")
    existing = attendance(date)
    if any(r["Id"] == str(sid) for r in existing):
        return False, "Attendance already marked today."
    try:
        start = datetime.datetime.strptime(SET["class_start_time"], "%H:%M").time()
    except Exception:
        start = datetime.time(9, 0)
    status = "Late" if datetime.datetime.now().time() > start else "On Time"
    p = attendance_file(date)
    if p.exists():
        wb = load_workbook(p); ws = wb.active
    else:
        wb = Workbook(); ws = wb.active
        ws.append(["Id", "Name", "Date", "In-Time", "Status"])
    ws.append([sid, name, date, datetime.datetime.now().strftime("%H:%M:%S"), status])
    wb.save(p); wb.close()
    return True, f"Attendance marked: {name} ({status})"

def recognize(data):
    if not MODEL.exists():
        return [], "Training model not found."
    if not hasattr(cv2, "face"):
        return [], "opencv-contrib-python is required."
    gray, fs = faces_from_bytes(data)
    if len(fs) == 0:
        return [], "No face detected."
    lookup = {r[0]: (r[1], r[2]) for r in students()}
    rec = cv2.face.LBPHFaceRecognizer_create()
    rec.read(str(MODEL))
    result = []
    for x, y, w, h in fs:
        serial, conf = rec.predict(gray[y:y+h, x:x+w])
        if serial in lookup and conf < float(SET["confidence_threshold"]):
            sid, name = lookup[serial]
            result.append((sid, name, round(float(conf), 2)))
    return result, None

st.set_page_config(page_title="FaceTrack Pro", page_icon="📷", layout="wide")
st.title("FaceTrack Pro")
st.caption("Smart Attendance & Face Recognition System")

if cascade() is None:
    st.error("Haar Cascade file is missing/invalid. Put haarcascade_frontalface_default.xml beside streamlit_app.py.")

t1,t2,t3,t4,t5,t6 = st.tabs(["Dashboard","Take Attendance","Students","Register","Reports","Settings"])

with t1:
    rows = students(); today = attendance()
    a,b,c = st.columns(3)
    a.metric("Total Registered", len(rows))
    b.metric("Present Today", len(today))
    c.metric("Late Today", sum(x["Status"].lower()=="late" for x in today))
    st.subheader("Today's Attendance")
    st.dataframe(pd.DataFrame(today), use_container_width=True, hide_index=True) if today else st.info("No attendance today.")

with t2:
    st.subheader("Take Attendance")
    photo = st.camera_input("Capture face")
    if photo:
        found, err = recognize(photo)
        if err: st.error(err)
        elif not found: st.warning("No registered student matched.")
        else:
            for sid,name,conf in found:
                st.success(f"Recognized: {name} | ID: {sid} | Confidence: {conf}")
                ok,msg = mark(sid,name)
                st.success(msg) if ok else st.info(msg)

with t3:
    st.subheader("Students")
    data = students()
    q = st.text_input("Search by ID or name")
    if q:
        data = [r for r in data if q.lower() in r[1].lower() or q.lower() in r[2].lower()]
    st.dataframe(pd.DataFrame(data, columns=["SERIAL NO.","ID","NAME"]), use_container_width=True, hide_index=True) if data else st.info("No students found.")

with t4:
    st.subheader("Register Student")
    sid = st.text_input("Student ID")
    name = st.text_input("Student Name")
    if st.button("Create Student Profile"):
        if not sid.strip() or not name.strip():
            st.error("Student ID and name are required.")
        elif any(r[1] == sid.strip() for r in students()):
            st.error("Student ID already exists.")
        else:
            serial = next_serial()
            wb = load_workbook(STUDENT_XLSX); wb.active.append([serial,sid.strip(),name.strip()]); wb.save(STUDENT_XLSX); wb.close()
            st.session_state["reg"] = (serial,sid.strip(),name.strip(),0)
            st.success(f"Profile created. Serial: {serial}")
    if "reg" in st.session_state:
        serial,sid,name,count = st.session_state["reg"]
        st.info(f"Capture samples for {name}. Samples: {count}")
        photo = st.camera_input("Training photo", key=f"training_{count}")
        if photo and count < 60:
            ok,msg = save_sample(serial,sid,name,photo,count+1)
            if ok:
                st.session_state["reg"] = (serial,sid,name,count+1)
                st.success(f"Sample {count+1} saved.")
            else: st.error(msg)
        if st.button("Train Face Model"):
            ok,msg = train()
            st.success(msg) if ok else st.error(msg)

with t5:
    st.subheader("Attendance Reports")
    today = datetime.date.today()
    start = st.date_input("Start Date", today)
    end = st.date_input("End Date", today)
    if start <= end:
        all_rows=[]; d=start
        while d<=end:
            all_rows += attendance(d.strftime("%d-%m-%Y")); d += datetime.timedelta(days=1)
        if all_rows:
            df=pd.DataFrame(all_rows); st.dataframe(df,use_container_width=True,hide_index=True)
            st.download_button("Download CSV", df.to_csv(index=False).encode(), "FaceTrackPro_Attendance.csv","text/csv")
        else: st.info("No records found.")
    else: st.error("Start Date cannot be after End Date.")

with t6:
    st.subheader("Settings")
    tm = st.time_input("Class Start Time", datetime.datetime.strptime(SET["class_start_time"],"%H:%M").time())
    threshold = st.slider("Recognition Confidence Threshold",10,100,int(SET["confidence_threshold"]))
    if st.button("Save Settings"):
        SET["class_start_time"]=tm.strftime("%H:%M"); SET["confidence_threshold"]=threshold
        SETTINGS_FILE.write_text(json.dumps(SET,indent=2),encoding="utf-8")
        st.success("Settings saved.")
